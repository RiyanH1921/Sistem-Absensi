import io
from flask import Blueprint, request, jsonify, send_file, g
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.services.report_service import ReportService
from app.models.academic import Class, Student, Subject
from app.models.attendance import Lesson, Attendance
from app.models.log import Notification
from app.utils.decorators import role_required, get_current_user

report_bp = Blueprint('report', __name__, url_prefix='/api/v1')

@report_bp.route('/dashboard/summary', methods=['GET'])
@role_required('superadmin', 'admin', 'teacher')
def dashboard_summary():
    data = ReportService.get_dashboard_summary()
    return jsonify({'success': True, 'data': data})

@report_bp.route('/reports/attendance', methods=['GET'])
@role_required('superadmin', 'admin', 'teacher')
def attendance_chart():
    class_id = request.args.get('class_id')
    month = request.args.get('month')
    year = request.args.get('year')
    data = ReportService.get_attendance_chart_data(class_id=class_id, month=month, year=year)
    return jsonify({'success': True, 'data': data})

@report_bp.route('/attendance/student/<int:student_id>', methods=['GET'])
def student_monthly_grid(student_id):
    subject_id = request.args.get('subject_id', type=int)
    month = request.args.get('month', type=int)
    year = request.args.get('year', type=int)

    data = ReportService.get_student_monthly_grid(student_id, subject_id, month, year)
    return jsonify({'success': True, 'data': data})

@report_bp.route('/notifications', methods=['GET'])
def get_notifications():
    user = get_current_user()
    if not user:
        return jsonify({'success': False, 'error': {'code': 'UNAUTHORIZED', 'message': 'Belum login'}}), 401

    notifs = Notification.query.filter_by(user_id=user.id).order_by(Notification.created_at.desc()).limit(20).all()
    unread_count = Notification.query.filter_by(user_id=user.id, is_read=False).count()

    return jsonify({
        'success': True,
        'unread_count': unread_count,
        'notifications': [n.to_dict() for n in notifs]
    })

@report_bp.route('/reports/export', methods=['GET'])
@role_required('superadmin', 'admin', 'teacher')
def export_excel():
    class_id = request.args.get('class_id', type=int)
    month = request.args.get('month', type=int)
    year = request.args.get('year', type=int)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Presensi"

    # Styling header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    headers = ["No", "NIS", "Nama Siswa", "Kelas", "Mata Pelajaran", "Tanggal", "Status", "Metode", "Waktu Scan", "Catatan"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    query = Attendance.query.join(Lesson).join(Student)
    if class_id:
        query = query.filter(Lesson.class_id == class_id)
    if month and year:
        query = query.filter(db.extract('month', Lesson.date) == month, db.extract('year', Lesson.date) == year)

    records = query.all()
    status_trans = {
        'present': 'Hadir',
        'late': 'Terlambat',
        'sick': 'Sakit',
        'permit': 'Izin',
        'absent': 'Alpa'
    }

    for idx, att in enumerate(records, 1):
        row = [
            idx,
            att.student.nis,
            att.student.full_name,
            att.lesson.class_room.name if att.lesson.class_room else '-',
            att.lesson.subject.name if att.lesson.subject else '-',
            att.lesson.date.strftime('%Y-%m-%d') if att.lesson.date else '-',
            status_trans.get(att.status, att.status.upper()),
            att.method.upper(),
            att.checked_in_at.strftime('%H:%M:%S') if att.checked_in_at else '-',
            att.note or '-'
        ]
        ws.append(row)

    # Auto adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Rekap_Presensi_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
